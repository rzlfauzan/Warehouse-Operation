import json
from pathlib import Path
ROOT=Path(__file__).resolve().parent.parent
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter as L
D=json.load(open(ROOT/'data/audit_dataset_jkt_v2.json'))
FLOWS=json.load(open(ROOT/'data/flows.json'))
AREA=FLOWS['areas']
def stage_rows(fl):
    """Ratakan data/flows.json menjadi daftar tahap (satu dict per tahap, urutan alur)."""
    out=[]
    for p in fl['procs']:
        for ph in p['phases']:
            n=0
            for st in ph['stages']:
                n+=1
                out.append(dict(flow=p['label'],phase=ph['name'],no=None if ph.get('periodic') else st['no'],name=st['name'],cond=st['cond'],
                    pic=st['pic'],main=AREA[st['areas'][0]],others=[AREA[a] for a in st['areas'][1:]],sop=st['sop'],
                    ref=[st['sheet'],st['from'],st['to']],**{'in':st['in'],'out':st['out']},sys=st['sys']))
    return out
rows=stage_rows(FLOWS)
F='Arial'
SHEETS={s['sheet']:s for s in D['sheets']}
PROC=[('Export','EXPORT'),('Import','IMPORT'),('Trucking Support','TRUCKING'),('DC','DC')]
FLOWN={'Export':'Export','Import':'Import','Trucking Support':'Trucking Support','DC':'DC (Distribution Center)'}
hf=Font(name=F,size=10,bold=True,color='FFFFFF'); navy=PatternFill('solid',fgColor='0B4A6F'); brown=PatternFill('solid',fgColor='7A5C00')
yel=PatternFill('solid',fgColor='FFF9C4'); stagef=PatternFill('solid',fgColor='E9EEF3')
areaf={'Gudang MKT':'DCE8F7','Port / Pelayaran':'D3F0F0','Depo':'FBE3D4','Bea Cukai':'F6D5E8'}
thin=Side(style='thin',color='CFD8E0'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='top')
HDR=['Tahap alur','Seksi audit','Area kerja','PIC tahap','No','Aktor','Kat','Acuan SOP','Langkah SOP / keterangan','Dokumen masuk','Dokumen keluar',
     'Catatan reviewer','Perlu diubah?','Usulan perubahan','Reviewer','Status review']
W=[26,9,17,24,6,20,7,17,50,32,32,34,12,32,13,15]
REV=12
def dl(ds):return '\n'.join('- '+d['t']+(' (TTD)' if d['sig'] else '') for d in ds)
def lines(v,w):return sum(max(1,-(-len(l)//max(1,int(w*1.05)))) for l in str(v).split('\n')) if v else 1

wb=Workbook(); wb.remove(wb.active)
counts={}
for name,code in PROC:
    ws=wb.create_sheet(name); ws.append(HDR)
    sh=SHEETS[code]; items=sorted(sh['items'],key=lambda i:i['no'])
    secname={s['id']:s['name'] for s in sh['sections']}
    stages=[r for r in rows if r['flow']==FLOWN[name]]
    covered=set(); nstage=0
    for st in stages:
        a,b=st['ref'][1],st['ref'][2]
        its=[i for i in items if a<=i['no']<=b]; covered|={i['no'] for i in its}
        label=(f"{st['no']}. " if st['no'] else 'Periodik: ')+st['name']+(f" ({st['cond']})" if st['cond'] else '')
        secs=sorted({i['section'] for i in its})
        area=st['main']+(' + '+', '.join(st['others']) if st['others'] else '')
        ws.append([label,', '.join(secs),area,', '.join(st['pic']),None,None,None,st['sop'],('Sistem / catatan: '+st['sys']) if st['sys'] else None,dl(st['in']),dl(st['out']),None,None,None,None,'Belum direview'])
        r=ws.max_row
        for c in range(1,len(HDR)+1):
            cell=ws.cell(r,c); cell.font=Font(name=F,size=10,bold=True if c<=4 else False); cell.fill=stagef
        ws.cell(r,3).fill=PatternFill('solid',fgColor=areaf.get(st['main'],'E9EEF3'))
        ws.cell(r,9).font=Font(name=F,size=10,italic=True)
        for i in its:
            k=i['criterion']+(f" [asal Export {i['orig_no']}]" if i.get('orig_no') else '')
            ws.append([label,i['section'],None,None,i['no'],i['actor'],i['category'],i['sop_ref'],k,None,None,None,None,None,None,'Belum direview'])
            ws.cell(ws.max_row,1).font=Font(name=F,size=9,color='7A8793')
    assert covered=={i['no'] for i in items},(name,{i['no'] for i in items}-covered)
    # style
    for c in range(1,len(HDR)+1):
        h=ws.cell(1,c);h.font=hf;h.fill=brown if c>=REV else navy;h.alignment=Alignment(wrap_text=True,vertical='center');h.border=bd
        ws.column_dimensions[L(c)].width=W[c-1]
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row,max_col=len(HDR)):
        mx=1
        for c in row:
            f=c.font; c.font=Font(name=F,size=f.size or 10,bold=f.bold,italic=f.italic,color=f.color)
            c.alignment=wrap;c.border=bd
            if c.column>=REV:c.fill=yel
            mx=max(mx,lines(c.value,W[c.column-1]))
        ws.row_dimensions[row[0].row].height=max(15,13.5*mx)
    ws.row_dimensions[1].height=30
    ws.freeze_panes='B2'; ws.auto_filter.ref=f"A1:{L(len(HDR))}{ws.max_row}"
    n=ws.max_row
    d1=DataValidation(type='list',formula1='"Ya,Tidak,Perlu cek"',allow_blank=True);ws.add_data_validation(d1);d1.add(f'M2:M{n}')
    d2=DataValidation(type='list',formula1='"Belum direview,Sedang direview,Sudah direview,Disetujui"',allow_blank=False);ws.add_data_validation(d2);d2.add(f'P2:P{n}')
    ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True);ws.print_title_rows='1:1'
    counts[name]=n

# Dokumen
wd=wb.create_sheet('Dokumen'); h2=['Proses','Tahap alur','Jenis','Dokumen','Perlu TTD','Catatan reviewer','Perlu diubah?','Nama resmi / usulan revisi','Reviewer']
wd.append(h2)
for x in rows:
    t=(f"{x['no']}. " if x['no'] else 'Periodik: ')+x['name']
    for j,lst in (('Masuk',x['in']),('Keluar',x['out'])):
        for d in lst: wd.append([x['flow'],t,j,d['t'],'Ya' if d['sig'] else '','','','',''])
W2=[22,36,9,50,10,40,13,36,14]
for c in range(1,len(h2)+1):
    h=wd.cell(1,c);h.font=hf;h.fill=brown if c>=6 else navy;h.alignment=Alignment(wrap_text=True,vertical='center');h.border=bd;wd.column_dimensions[L(c)].width=W2[c-1]
for row in wd.iter_rows(min_row=2,max_row=wd.max_row,max_col=len(h2)):
    mx=1
    for c in row:
        c.font=Font(name=F,size=10);c.alignment=wrap;c.border=bd
        if c.column>=6:c.fill=yel
        mx=max(mx,lines(c.value,W2[c.column-1]))
    wd.row_dimensions[row[0].row].height=max(15,13.5*mx)
wd.freeze_panes='C2';wd.auto_filter.ref=f"A1:{L(len(h2))}{wd.max_row}"
dv=DataValidation(type='list',formula1='"Ya,Tidak,Perlu cek"',allow_blank=True);wd.add_data_validation(dv);dv.add(f'G2:G{wd.max_row}')

# Petunjuk
wp=wb.create_sheet('Petunjuk',0); wp.sheet_view.showGridLines=False
for c,w in zip('ABCD',[30,16,26,50]):wp.column_dimensions[c].width=w
def put(r,c,v,**k):
    x=wp.cell(r,c,v);x.font=Font(name=F,size=k.get('size',10),bold=k.get('bold',False),color=k.get('color','000000'));x.alignment=Alignment(wrap_text=True,vertical='top')
    if k.get('fill'):x.fill=k['fill']
    return x
put(1,1,'Peta Proses Audit SOP: lembar review',size=14,bold=True)
put(2,1,'Sumber: form audit SOP FRM-OPS-001-0001 (Jakarta, Agustus 2026), untuk SOP Warehouse dan Multimode MKT. Isi tahap dan dokumen sama dengan halaman Alur Proses Umum MKT. Hasil audit tidak dimuat.',color='5A6B78')
wp.merge_cells('A2:D2'); wp.row_dimensions[2].height=30
put(4,1,'Cara pakai',bold=True)
steps=['Satu sheet per proses: Export, Import, Trucking Support, DC. Tiap tahap alur diawali satu baris berlatar abu (area kerja, PIC, dokumen masuk dan keluar), diikuti langkah audit di bawahnya.',
 'Dokumen masuk dan keluar dimuat per tahap, bukan per langkah, karena dokumennya milik tahap. Langkah audit (No, Aktor, Acuan SOP) mengikuti nomor di Peta Proses Audit SOP.',
 'Sheet "Dokumen" memuat satu baris per dokumen untuk catatan per dokumen (nama resmi, perlu TTD).',
 'Kolom berwarna kuning adalah kolom review dan satu-satunya yang perlu diisi. "Perlu diubah?" dan "Status review" memakai dropdown.',
 '(TTD) berarti dokumen perlu tanda tangan. "(perkiraan)" berarti penempatan dokumen di tahap itu belum dipastikan.',
 'Trucking Support langkah 1 sampai 8 (penarikan empty) asalnya Export 37 sampai 44. Langkah 9 sampai 22 (penarikan full container) belum diaudit dan masih draf dari Pendahuluan SOP.']
for i,s in enumerate(steps):
    put(5+i,1,f'{i+1}. '+s);wp.merge_cells(start_row=5+i,start_column=1,end_row=5+i,end_column=4);wp.row_dimensions[5+i].height=30
r=12
put(r,1,'Contoh pengisian (hanya contoh, tidak ada di sheet proses)',bold=True)
for i,(k,v) in enumerate([('Catatan reviewer','Tahap ini di lapangan dikerjakan Admin Warehouse, bukan AO.'),('Perlu diubah?','Ya'),('Usulan perubahan','Ubah PIC menjadi Admin Warehouse'),('Status review','Sudah direview')]):
    put(r+1+i,1,k,fill=stagef);put(r+1+i,2,v,fill=yel);wp.merge_cells(start_row=r+1+i,start_column=2,end_row=r+1+i,end_column=4)
r=18
put(r,1,'Progres review (otomatis)',bold=True)
for c,t in enumerate(['Sheet','Jumlah baris','Sudah direview atau disetujui','Ditandai perlu diubah'],1):put(r+1,c,t,bold=True,color='FFFFFF',fill=navy)
for i,(name,code) in enumerate(PROC):
    rr=r+2+i;n=counts[name];q=f"'{name}'"
    put(rr,1,name)
    wp.cell(rr,2,f"=COUNTA({q}!$A$2:$A${n})")
    wp.cell(rr,3,f'=COUNTIF({q}!$P$2:$P${n},"Sudah direview")+COUNTIF({q}!$P$2:$P${n},"Disetujui")')
    wp.cell(rr,4,f'=COUNTIF({q}!$M$2:$M${n},"Ya")')
    for c in (2,3,4):wp.cell(rr,c).font=Font(name=F,size=10)
tr=r+2+len(PROC);put(tr,1,'Total',bold=True)
for c in (2,3,4):
    x=wp.cell(tr,c,f'=SUM({L(c)}{r+2}:{L(c)}{tr-1})');x.font=Font(name=F,size=10,bold=True)
wb.save(ROOT/'dist/Peta_Proses_Audit_SOP_Review.xlsx'); print(counts)
