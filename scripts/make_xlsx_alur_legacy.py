"""Catatan: skrip sejarah untuk Excel Alur Proses Umum yang kemudian diganti make_xlsx_peta.py. Path merujuk folder sementara."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
rows=json.load(open('flow_rows.json'))
F='Arial'
hf=Font(name=F,size=10,bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='0B4A6F')
rf=PatternFill('solid',fgColor='FFF9C4'); gf=PatternFill('solid',fgColor='E9EEF3')
thin=Side(style='thin',color='CFD8E0'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='top')
def style(ws,hdr_row,ncols,review_from,widths):
    for c in range(1,ncols+1):
        h=ws.cell(hdr_row,c);h.font=hf;h.fill=PatternFill('solid',fgColor='7A5C00' if c>=review_from else '0B4A6F')
        h.alignment=Alignment(wrap_text=True,vertical='center');h.border=bd
        ws.column_dimensions[L(c)].width=widths[c-1]
    for r in ws.iter_rows(min_row=hdr_row+1,max_row=ws.max_row,max_col=ncols):
        mx=1
        for c in r:
            c.font=Font(name=F,size=10);c.alignment=wrap;c.border=bd
            if c.column>=review_from:c.fill=rf
            w=widths[c.column-1]
            if c.value:mx=max(mx,sum(max(1,-(-len(l)//max(1,int(w*1.05)))) for l in str(c.value).split('\n')))
        ws.row_dimensions[r[0].row].height=max(15,13.5*mx)
    ws.freeze_panes=ws.cell(hdr_row+1,5 if ncols>12 else 4)
    ws.auto_filter.ref=f"A{hdr_row}:{L(ncols)}{ws.max_row}"
    ws.row_dimensions[hdr_row].height=32

wb=Workbook()
# ---------- Tahap
ws=wb.active;ws.title='Tahap'
hdr=['Proses','Fase','No','Tahap','Kondisional','Area utama','Area terlibat','PIC','Dokumen masuk','Dokumen keluar','Sistem / catatan','Acuan SOP','Langkah audit','Hasil audit (as reported)',
     'Catatan reviewer','Perlu diubah?','Usulan perubahan','Reviewer','Status review']
ws.append(hdr)
def audit_ref(x):
    sh,a,b=x['ref']
    nm={'EXPORT':'Export','IMPORT':'Import','DC':'DC','TRUCKING':'Trucking'}[sh]
    t=f"{nm} {a}" if a==b else f"{nm} {a}-{b}"
    if sh=='TRUCKING' and b<=8: t+=f" (asal Export {a+36}-{b+36})" if a!=b else f" (asal Export {a+36})"
    return t
def audit_res(x):
    if x['pending']:return 'Belum diaudit'
    if x['flow'].startswith('DC'):
        s,m,d=x['nc'];return f"NC: Semper {s} | Marcen {m} | Dokumen {d}"
    return f"NC: {x['nc']}"
def dl(ds):return '\n'.join(('- '+d['t']+(' (TTD)' if d['sig'] else '')) for d in ds)
for x in rows:
    ws.append([x['flow'],x['phase'],x['no'] if x['no'] else 'Periodik',x['name'],x['cond'],x['main'],', '.join(x['others']),', '.join(x['pic']),
      dl(x['in']),dl(x['out']),x['sys'],x['sop'],audit_ref(x),audit_res(x),'','','','','Belum direview'])
style(ws,1,len(hdr),15,[16,20,9,26,11,15,16,24,34,34,34,16,22,22,36,13,36,14,15])
n=ws.max_row
dv=DataValidation(type='list',formula1='"Ya,Tidak,Perlu cek"',allow_blank=True);ws.add_data_validation(dv);dv.add(f'P2:P{n}')
dv2=DataValidation(type='list',formula1='"Belum direview,Sedang direview,Sudah direview,Disetujui"',allow_blank=False);ws.add_data_validation(dv2);dv2.add(f'S2:S{n}')

# ---------- Dokumen
wd=wb.create_sheet('Dokumen')
h2=['Proses','Tahap','Jenis','Dokumen','Perlu TTD','Catatan reviewer','Perlu diubah?','Nama resmi / usulan revisi','Reviewer']
wd.append(h2)
for x in rows:
    t=f"{x['no']}. {x['name']}" if x['no'] else f"Periodik: {x['name']}"
    for j,lst in (('Masuk',x['in']),('Keluar',x['out'])):
        for d in lst:wd.append([x['flow'],t,j,d['t'],'Ya' if d['sig'] else '','','','',''])
style(wd,1,len(h2),6,[18,34,9,48,10,40,13,36,14])
m=wd.max_row
dv3=DataValidation(type='list',formula1='"Ya,Tidak,Perlu cek"',allow_blank=True);wd.add_data_validation(dv3);dv3.add(f'G2:G{m}')

# ---------- Petunjuk
wp=wb.create_sheet('Petunjuk',0)
wp.sheet_view.showGridLines=False
wp.column_dimensions['A'].width=30;wp.column_dimensions['B'].width=16;wp.column_dimensions['C'].width=70
def put(r,c,v,**k):
    x=wp.cell(r,c,v);x.font=Font(name=F,size=k.get('size',10),bold=k.get('bold',False),color=k.get('color','000000'));x.alignment=Alignment(wrap_text=True,vertical='top')
    if k.get('fill'):x.fill=k['fill']
    return x
put(1,1,'Alur Proses Umum MKT: lembar review',size=14,bold=True)
put(2,1,'Sumber: form audit SOP FRM-OPS-001-0001 (Jakarta, Agustus 2026). Isi sama dengan halaman Alur Proses Umum MKT di Claude.',color='5A6B78')
wp.merge_cells('A2:C2')
put(4,1,'Cara pakai',bold=True)
steps=['Sheet "Tahap": satu baris per tahap proses (area kerja, PIC, dokumen masuk dan keluar, hasil audit).',
 'Sheet "Dokumen": satu baris per dokumen, untuk catatan per dokumen (nama resmi, perlu TTD, dan sebagainya).',
 'Kolom berwarna kuning adalah kolom review. Hanya kolom itu yang perlu diisi. Kolom lain berasal dari halaman dan dibiarkan agar bisa dicocokkan.',
 '"Perlu diubah?" dan "Status review" memakai pilihan dropdown. Tabel di bawah menghitung status review secara otomatis.',
 'Tanda (TTD) berarti dokumen memerlukan tanda tangan. "(perkiraan)" berarti penempatan dokumen di tahap itu belum dipastikan.',
 'Tahap Penarikan full container ke UTPK/UTC berstatus "Belum diaudit". Isinya masih draf dari Pendahuluan SOP, butir Prosedur belum diterima.']
for i,s in enumerate(steps):
    put(5+i,1,f'{i+1}. '+s);wp.merge_cells(start_row=5+i,start_column=1,end_row=5+i,end_column=3);wp.row_dimensions[5+i].height=28
r=12
put(r,1,'Contoh pengisian (hanya contoh, tidak ada di sheet Tahap)',bold=True)
ex=[('Catatan reviewer','Tahap ini di lapangan dikerjakan Admin Warehouse, bukan AO.'),('Perlu diubah?','Ya'),('Usulan perubahan','Ubah PIC menjadi Admin Warehouse'),('Status review','Sudah direview')]
for i,(k,v) in enumerate(ex):
    put(r+1+i,1,k,fill=gf);put(r+1+i,2,v,fill=rf);wp.merge_cells(start_row=r+1+i,start_column=2,end_row=r+1+i,end_column=3)
r=18
put(r,1,'Progres review (otomatis)',bold=True)
for c,t in enumerate(['Proses','Jumlah tahap','Sudah direview atau disetujui'],1):
    x=put(r+1,c,t,bold=True,color='FFFFFF',fill=hfill)
flows=list(dict.fromkeys(x['flow'] for x in rows))
for i,f in enumerate(flows):
    rr=r+2+i
    put(rr,1,f)
    wp.cell(rr,2,f'=COUNTIF(Tahap!$A$2:$A${n},A{rr})').font=Font(name=F,size=10)
    wp.cell(rr,3,f'=COUNTIFS(Tahap!$A$2:$A${n},A{rr},Tahap!$S$2:$S${n},"Sudah direview")+COUNTIFS(Tahap!$A$2:$A${n},A{rr},Tahap!$S$2:$S${n},"Disetujui")').font=Font(name=F,size=10)
tr=r+2+len(flows)
put(tr,1,'Total',bold=True)
wp.cell(tr,2,f'=SUM(B{r+2}:B{tr-1})').font=Font(name=F,size=10,bold=True)
wp.cell(tr,3,f'=SUM(C{r+2}:C{tr-1})').font=Font(name=F,size=10,bold=True)
put(tr+1,1,'Tahap yang ditandai "Perlu diubah? = Ya"');wp.cell(tr+1,2,f'=COUNTIF(Tahap!$P$2:$P${n},"Ya")').font=Font(name=F,size=10)
wb.save('/home/user/Alur_Proses_Umum_MKT_Review.xlsx')
print(n,m)
